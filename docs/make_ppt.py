#!/usr/bin/env python3
"""Build a 10-slide PPTX for the vision-only UAV navigation system using real project assets."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

REPO = Path(__file__).resolve().parents[1]
A = REPO / "docs" / "ppt_assets"
IMG = REPO / "docs" / "images"
KF = REPO / "linkedin_assets" / "vision_guided_keyframes_final"
LA = REPO / "linkedin_assets"

INK = RGBColor(0x16, 0x1B, 0x22)
MUT = RGBColor(0x55, 0x5F, 0x6B)
ACC = RGBColor(0x2B, 0x6C, 0xB0)      # blue
RED = RGBColor(0xB0, 0x2B, 0x2B)
GRN = RGBColor(0x1F, 0x7A, 0x44)
BG = RGBColor(0xFF, 0xFF, 0xFF)
BAND = RGBColor(0xF2, 0xF4, 0xF7)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(1, 0, 0, SW, SH); bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def rect(s, x, y, w, h, color):
    r = s.shapes.add_shape(1, x, y, w, h); r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background(); r.shadow.inherit = False
    return r


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, (t, sz, col, bold) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(4)
        r = p.add_run(); r.text = t
        r.font.size = Pt(sz); r.font.color.rgb = col; r.font.bold = bold
        r.font.name = "Segoe UI"
    return tb


def bullets(s, x, y, w, h, items, sz=16, col=INK, gap=7):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = "•  " + it
        r.font.size = Pt(sz); r.font.color.rgb = col; r.font.name = "Segoe UI"
    return tb


def header(s, n, kicker, title, accent=ACC):
    rect(s, 0, 0, Inches(0.18), SH, accent)
    text(s, Inches(0.6), Inches(0.35), Inches(11), Inches(0.4),
         [(kicker.upper(), 13, accent, True)])
    text(s, Inches(0.6), Inches(0.7), Inches(12), Inches(0.9), [(title, 30, INK, True)])
    text(s, Inches(12.2), Inches(7.0), Inches(1), Inches(0.4), [(f"{n}/10", 11, MUT, False)])


def pic(s, path, x, y, w, h=None):
    if not Path(path).exists():
        return
    if h is None:
        s.shapes.add_picture(str(path), x, y, width=w)
    else:
        s.shapes.add_picture(str(path), x, y, width=w, height=h)


def first(*cands):
    for c in cands:
        if Path(c).exists():
            return str(c)
    return None


# ---------- 1 title ----------
s = slide()
rect(s, 0, 0, SW, SH, INK)
hero = first(LA / "generated_hero_frame.png", KF / "keyframe_01.png")
if hero:
    s.shapes.add_picture(hero, Inches(7.0), Inches(1.1), height=Inches(5.3))
rect(s, 0, 0, Inches(0.22), SH, ACC)
text(s, Inches(0.7), Inches(1.4), Inches(6.4), Inches(0.5),
     [("VISION-ONLY UAV SYSTEM", 15, RGBColor(0x8F, 0xC2, 0xFF), True)])
text(s, Inches(0.7), Inches(1.9), Inches(6.6), Inches(2.2),
     [("Vision-only navigation & target tracking", 40, BG, True)])
text(s, Inches(0.7), Inches(4.2), Inches(6.2), Inches(1.4),
     [("Detect, track, self-localize and avoid — even when GPS is jammed.", 18,
       RGBColor(0xCF, 0xD8, 0xE3), False)])
text(s, Inches(0.7), Inches(6.4), Inches(8), Inches(0.5),
     [("Sim-validated in AirSim  ·  roadmap to real hardware (RealSense + Jetson)", 13,
       RGBColor(0x9A, 0xA6, 0xB4), False)])

# ---------- 2 problem ----------
s = slide(); header(s, 2, "the problem", "Why vision, not GPS", RED)
bullets(s, Inches(0.7), Inches(2.0), Inches(7.2), Inches(4.5), [
    "GPS can be jammed or spoofed in contested airspace.",
    "A GPS-reliant drone is effectively blind without it.",
    "Our approach: navigate by what the camera sees.",
    "One camera does two jobs — see the target, and know its own position.",
    "Vision is the backup that cannot be jammed.",
], sz=18, gap=12)
rect(s, Inches(8.3), Inches(2.0), Inches(4.3), Inches(3.6), BAND)
text(s, Inches(8.55), Inches(2.25), Inches(3.9), Inches(0.6), [("GPS vs VISION", 14, MUT, True)])
text(s, Inches(8.55), Inches(2.8), Inches(3.9), Inches(2.7), [
    ("GPS  →  jammable, single point of failure", 15, RED, False),
    ("", 6, MUT, False),
    ("VISION  →  passive, self-contained, jam-proof", 15, GRN, True),
], anchor=MSO_ANCHOR.TOP)

# ---------- 3 concept pipeline ----------
s = slide(); header(s, 3, "system concept", "How it works, end to end")
steps = [("SEE", "detect + segment", ACC), ("TRACK", "lock + predict", GRN),
         ("LOCATE", "VIO self-position", ACC), ("AVOID", "depth steering", RGBColor(0xB5, 0x7A, 0x12)),
         ("ACT", "follow + fly", GRN)]
x = Inches(0.7); w = Inches(2.18); gap = Inches(0.2)
for i, (t, sub, c) in enumerate(steps):
    bx = Emu(int(x) + i * (int(w) + int(gap)))
    rect(s, bx, Inches(2.6), w, Inches(1.6), BAND)
    rect(s, bx, Inches(2.6), w, Inches(0.12), c)
    text(s, bx, Inches(2.95), w, Inches(0.6), [(t, 19, INK, True)], align=PP_ALIGN.CENTER)
    text(s, bx, Inches(3.55), w, Inches(0.5), [(sub, 13, MUT, False)], align=PP_ALIGN.CENTER)
text(s, Inches(0.7), Inches(4.7), Inches(12), Inches(1.2), [
    ("One camera feed splits into two pipelines: tracking OTHERS (detect + follow the target) and "
     "locating SELF (visual-inertial odometry). They share pixels, run independently.", 16, MUT, False)])

# ---------- 4 perception real detections ----------
s = slide(); header(s, 4, "perception", "Seeing the target — on real data")
bullets(s, Inches(0.7), Inches(1.9), Inches(4.6), Inches(4.5), [
    "YOLO26 detector + instance segmentation.",
    "Two classes: drone target + bird (false-lock filter).",
    "Trained on ~56k real + synthetic images (Shahed, Mjolnir, AirSim).",
    "Real held-out mAP50 ≈ 0.77 and climbing.",
    "Boxes below are our model on unseen real frames.",
], sz=15, gap=9)
dets = sorted(A.glob("det_real_*.jpg"))[:4]
gx, gy = Inches(5.6), Inches(1.9); cw = Inches(3.5); ch = Inches(2.0)
for i, d in enumerate(dets):
    bx = Emu(int(gx) + (i % 2) * (int(cw) + Inches(0.15)))
    by = Emu(int(gy) + (i // 2) * (int(ch) + Inches(0.2)))
    pic(s, d, bx, by, cw, ch)

# ---------- 5 trained & tested ----------
s = slide(); header(s, 5, "training & testing", "How it learns — and how we prove it", GRN)
bullets(s, Inches(0.7), Inches(1.9), Inches(7.0), Inches(4.6), [
    "Learns from real + simulated drone footage.",
    "Segmentation outlines the exact target; filtering drops birds, clutter, shadows.",
    "Guidance is learned by imitating the classical expert, then improved with reinforcement learning (hybrid).",
    "Tested across orbit, zigzag, weave, climb and dive maneuvers in simulation.",
    "Imitation policy keeps the target in frame ~100% (sim); RL tightens the stand-off hold.",
], sz=16, gap=11)
kf = first(KF / "keyframe_03.png", IMG / "track_b.jpg")
pic(s, kf, Inches(8.1), Inches(2.1), Inches(4.5))
text(s, Inches(8.1), Inches(5.6), Inches(4.5), Inches(0.5),
     [("data → segmentation → filter → imitate → RL → sim test", 12, MUT, False)], align=PP_ALIGN.CENTER)

# ---------- 6 tracking / web UI ----------
s = slide(); header(s, 6, "tracking", "Target following — live web UI")
bullets(s, Inches(0.7), Inches(1.9), Inches(4.5), Inches(4.5), [
    "Click-to-lock the target in the live web UI.",
    "Predicts motion (Kalman), coasts through detection drop-outs.",
    "Holds a set stand-off distance, no spin.",
    "Kept in frame ~90%+ vision-only.",
], sz=15, gap=10)
tk = [IMG / "track_a.jpg", IMG / "track_b.jpg", IMG / "track_c.jpg"]
for i, t in enumerate(tk):
    pic(s, t, Inches(5.5), Emu(int(Inches(1.85)) + i * int(Inches(1.65))), Inches(7.0))

# ---------- 7 VIO ----------
s = slide(); header(s, 7, "self-localization", "VIO / VINS · depth · SLAM")
bullets(s, Inches(0.7), Inches(1.9), Inches(7.2), Inches(4.6), [
    "Camera + IMU estimate the drone's own position with no GPS (visual-inertial odometry).",
    "Depth camera gives true distance and a live 3D sense of the scene.",
    "Place recognition (SLAM-style) corrects accumulated drift.",
    "Magnetometer bounds heading drift.",
    "Production path: OpenVINS (MSCKF) — target under 1 m drift, fed to PX4 EKF2.",
], sz=16, gap=11)
kf = first(KF / "keyframe_05.png", KF / "keyframe_02.png")
pic(s, kf, Inches(8.2), Inches(2.1), Inches(4.4))

# ---------- 8 obstacle avoidance ----------
s = slide(); header(s, 8, "safety", "Obstacle avoidance", RGBColor(0xB5, 0x7A, 0x12))
bullets(s, Inches(0.7), Inches(1.9), Inches(7.2), Inches(4.5), [
    "Depth is read like a wall of distances (depth-as-LiDAR).",
    "The drone steers around whatever is close.",
    "It keeps chasing and framing the target while it dodges.",
    "Reactive — runs onboard, no map needed.",
], sz=17, gap=12)
kf = first(KF / "keyframe_06.png", KF / "keyframe_04.png")
pic(s, kf, Inches(8.2), Inches(2.1), Inches(4.4))

# ---------- 9 applications ----------
s = slide(); header(s, 9, "applications", "Where it's used")
rect(s, Inches(0.7), Inches(1.9), Inches(5.85), Inches(4.7), BAND)
rect(s, Inches(0.7), Inches(1.9), Inches(5.85), Inches(0.13), GRN)
text(s, Inches(0.95), Inches(2.1), Inches(5.4), Inches(0.5), [("COMMERCIAL", 16, GRN, True)])
bullets(s, Inches(0.95), Inches(2.7), Inches(5.4), Inches(3.7), [
    "Cinematography / follow-me filming",
    "Infrastructure inspection in GPS-poor sites",
    "Search & rescue in canyons / indoors",
    "Warehouse & indoor inventory drones",
    "Agriculture & survey, delivery in urban canyons",
], sz=14, gap=8)
rect(s, Inches(6.75), Inches(1.9), Inches(5.85), Inches(4.7), BAND)
rect(s, Inches(6.75), Inches(1.9), Inches(5.85), Inches(0.13), RED)
text(s, Inches(7.0), Inches(2.1), Inches(5.4), Inches(0.5), [("DEFENSE", 16, RED, True)])
bullets(s, Inches(7.0), Inches(2.7), Inches(5.4), Inches(3.7), [
    "Counter-UAS / anti-Shahed interception",
    "ISR in GPS-denied / jammed airspace",
    "Convoy overwatch & perimeter security",
    "Contested-airspace navigation & strike",
    "Comms-denied autonomous missions",
], sz=14, gap=8)

# ---------- 10 summary ----------
s = slide(); header(s, 10, "summary", "Status & roadmap")
bullets(s, Inches(0.7), Inches(1.9), Inches(11.8), Inches(4.0), [
    "Vision-only detect → track → follow validated in AirSim simulation.",
    "Detector now trained on real air-to-air data (Shahed / Mjolnir) for real-world transfer.",
    "Guidance policy: imitation of the classical expert, then RL fine-tune (hybrid).",
    "Sim-to-real: the SAME software, swap the sensor — RealSense D455 + Jetson Orin.",
    "Next: OpenVINS VIO → PX4 EKF2 for GPS-denied flight; lightweight agentic supervisor last.",
], sz=17, gap=12)
rect(s, Inches(0.7), Inches(6.0), Inches(11.9), Inches(0.9), INK)
text(s, Inches(0.7), Inches(6.05), Inches(11.9), Inches(0.8),
     [("The eye is the backup that can't be jammed.", 20, BG, True)],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

out = REPO / "docs" / "UAV_Vision_Navigation.pptx"
prs.save(str(out))
print("saved", out, "slides", len(prs.slides._sldIdLst))
